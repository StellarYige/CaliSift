"""Rebuild checked-in, synthetic Excel/CSV fixtures. No personal data."""

from __future__ import annotations

import csv
import io
from datetime import datetime, time
from pathlib import Path

import openpyxl
import xlwt
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
NAME = "星辰奕歌"


def xlsx_bytes(
    sheets: dict[str, list[list]], merges=None, hidden=None, epoch=None
) -> bytes:
    book = openpyxl.Workbook()
    book.remove(book.active)
    if epoch:
        book.epoch = epoch
    for name, rows in sheets.items():
        ws = book.create_sheet(name)
        for row in rows:
            ws.append(row)
        for span in (merges or {}).get(name, []):
            ws.merge_cells(span)
        for row in ws:
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if isinstance(cell.value, datetime):
                    cell.number_format = (
                        "yyyy-mm-dd hh:mm"
                        if cell.value.time() != time()
                        else "yyyy-mm-dd"
                    )
                elif isinstance(cell.value, time):
                    cell.number_format = "hh:mm"
        for cell in ws[1]:
            cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="39406B")
        for col in ws.columns:
            ws.column_dimensions[
                (
                    col[0].column_letter
                    if not isinstance(col[0], openpyxl.cell.cell.MergedCell)
                    else openpyxl.utils.get_column_letter(col[0].column)
                )
            ].width = 25
        for r in range(1, ws.max_row + 1):
            ws.row_dimensions[r].height = 34
        ws.freeze_panes = "B2"
        if name in (hidden or []):
            ws.sheet_state = "hidden"
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def xls_bytes(rows, merges=None, datemode=False) -> bytes:
    book = xlwt.Workbook()
    book.set_dates_1904(datemode)
    ws = book.add_sheet("值班")
    skip = set()
    for r0, r1, c0, c1 in merges or []:
        ws.write_merge(r0, r1, c0, c1, rows[r0][c0])
        skip.update((r, c) for r in range(r0, r1 + 1) for c in range(c0, c1 + 1))
    for r, row in enumerate(rows):
        ws.row(r).height = 480
        for c, value in enumerate(row):
            ws.col(c).width = 6500
            if (r, c) not in skip and value is not None:
                style = (
                    xlwt.easyxf(num_format_str="YYYY-MM-DD")
                    if isinstance(value, datetime)
                    else xlwt.Style.default_style
                )
                ws.write(r, c, value, style)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def csv_bytes(rows, encoding="utf-8-sig", delimiter=",") -> bytes:
    output = io.StringIO(newline="")
    csv.writer(output, delimiter=delimiter).writerows(rows)
    return output.getvalue().encode(encoding)


def generate(destination: Path):
    destination.mkdir(parents=True, exist_ok=True)
    fixtures = {
        "9月排班表.xlsx": xlsx_bytes(
            {
                "九月": [
                    ["2026年9月排班表", None, None],
                    ["姓名", "9月7日", "9月8日"],
                    [NAME, "早班", "休息"],
                    ["林知夏", "晚班", "早班"],
                    [],
                    ["早班：08:00–16:00"],
                    ["晚班：16:00–23:00"],
                ]
            },
            {"九月": ["A1:C1"]},
        ),
        "培训安排.xlsx": xlsx_bytes(
            {
                "消防": [
                    ["日期", "时间", "事项", "参加人员", "地点", "备注"],
                    [
                        "2026年9月12日",
                        "下午2点",
                        "消防培训",
                        f"林知夏、{NAME}",
                        "一楼培训室",
                        "提前十分钟到场",
                    ],
                ],
                "其他培训": [
                    ["日期", "事项", "参加人员"],
                    ["2026年9月13日", "设备培训", "林知夏"],
                ],
            }
        ),
        "国庆值班表.xlsx": xlsx_bytes(
            {
                "国庆": [
                    ["日期", "事项", "值班人员", "地点"],
                    ["2026年10月2日", "国庆值班", NAME, "服务中心"],
                ]
            }
        ),
        "冲突与重复.xlsx": xlsx_bytes(
            {
                "核对": [
                    ["日期", "时间", "班次", "事项", "姓名", "备注"],
                    [
                        "2026-09-07",
                        "08:00-16:00",
                        "早班",
                        "早班",
                        NAME,
                        "与九月排班重复",
                    ],
                    ["2026-09-07", "15:00-17:00", "", "应急演练", NAME, "与早班冲突"],
                ]
            }
        ),
        "纵向排班.xls": xls_bytes(
            [
                ["2026年9月值班表", None, None],
                ["日期", NAME, "林知夏"],
                [datetime(2026, 9, 9), "夜班 22:00-次日06:00", "早班"],
                [datetime(2026, 9, 10), "休息", "早班"],
            ],
            [(0, 0, 0, 2)],
        ),
        "考试安排.csv": csv_bytes(
            [
                ["日期", "开始时间", "结束时间", "考试科目", "监考人员", "考场"],
                ["2026年9月15日", "09:00", "11:00", "计算机基础", NAME, "301"],
            ],
            "gb18030",
        ),
        "合并表头.xlsx": xlsx_bytes(
            {
                "2026年9月": [
                    ["日期", "9月7日", None, "9月8日", None],
                    ["班次", "早班", "晚班", "早班", "晚班"],
                    ["人员", f"{NAME}、林知夏", "陈清和", "林知夏", NAME],
                ]
            },
            {"2026年9月": ["B1:C1", "D1:E1"]},
        ),
        "周期课表待确认.xlsx": xlsx_bytes(
            {
                "课表": [
                    ["星期", "节次", "课程", "教师"],
                    ["星期一", "第1–2节", "数学", NAME],
                ]
            }
        ),
    }
    for filename, data in fixtures.items():
        (destination / filename).write_bytes(data)
    return list(fixtures)


if __name__ == "__main__":
    for filename in generate(ROOT / "tests" / "fixtures"):
        print(filename)
