from __future__ import annotations

import csv
import hashlib
import io
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
import xlrd
from openpyxl.utils import get_column_letter

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_CELLS = 250_000
MAX_SHEETS = 50
MAX_EXPANDED_BYTES = 60 * 1024 * 1024


class FileProblem(ValueError):
    pass


def text(value: Any) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip()


@dataclass
class Cell:
    row: int
    col: int
    value: Any = None
    formula_missing: bool = False
    anchor: tuple[int, int] | None = None

    @property
    def coordinate(self):
        r, c = self.anchor or (self.row, self.col)
        return f"{get_column_letter(c + 1)}{r + 1}"

    @property
    def label(self):
        return text(self.value)


@dataclass
class Sheet:
    name: str
    rows: list[list[Cell]]
    hidden: bool = False
    warnings: list[str] = field(default_factory=list)

    @property
    def height(self):
        return len(self.rows)

    @property
    def width(self):
        return len(self.rows[0]) if self.rows else 0

    def at(self, r: int, c: int) -> Cell:
        if 0 <= r < self.height and 0 <= c < self.width:
            return self.rows[r][c]
        return Cell(r, c)


def _shape(rows: int, cols: int, total: int) -> int:
    total += rows * cols
    if total > MAX_CELLS:
        raise FileProblem(f"工作簿超过 {MAX_CELLS:,} 个单元格的处理上限")
    return total


def _merged(sheet: Sheet, ranges):
    for r0, r1, c0, c1 in ranges:
        anchor = sheet.at(r0, c0)
        if r1 > sheet.height or c1 > sheet.width:
            raise FileProblem("合并单元格范围超出工作表边界")
        for r in range(r0, r1):
            for c in range(c0, c1):
                cell = sheet.at(r, c)
                cell.value = anchor.value
                cell.formula_missing = anchor.formula_missing
                cell.anchor = (r0, c0)


def read_workbook(data: bytes, filename: str) -> tuple[str, list[Sheet]]:
    if not data:
        raise FileProblem("文件为空")
    if len(data) > MAX_FILE_BYTES:
        raise FileProblem("单个文件不能超过 10 MiB")
    ext = Path(filename).suffix.lower()
    file_id = hashlib.sha256(data).hexdigest()[:20]
    try:
        if ext == ".xlsx":
            sheets = _xlsx(data)
        elif ext == ".xls":
            sheets = _xls(data)
        elif ext == ".csv":
            sheets = _csv(data)
        else:
            raise FileProblem("仅支持 .xlsx、.xls、.csv 文件")
    except FileProblem:
        raise
    except Exception as exc:
        raise FileProblem(
            "文件损坏、已加密或内容与扩展名不符，请另存为 Excel/CSV 后重试"
        ) from exc
    return file_id, sheets


def _xlsx(data: bytes) -> list[Sheet]:
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise FileProblem("不是有效的 xlsx 文件，可能已加密或扩展名不正确")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if (
            len(entries) > 5000
            or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES
        ):
            raise FileProblem("xlsx 解压后过大，无法处理")
        if any(e.flag_bits & 1 for e in entries):
            raise FileProblem("暂不支持加密工作簿")
    formulas = openpyxl.load_workbook(
        io.BytesIO(data), data_only=False, keep_links=False
    )
    values = None
    try:
        if len(formulas.worksheets) > MAX_SHEETS:
            raise FileProblem("工作表数量不能超过 50")
        values = openpyxl.load_workbook(
            io.BytesIO(data), data_only=True, keep_links=False
        )
        result, total = [], 0
        for ws in formulas.worksheets:
            total = _shape(ws.max_row, ws.max_column, total)
            cached = values[ws.title]
            rows = []
            missing = []
            for r in range(ws.max_row):
                row = []
                for c in range(ws.max_column):
                    cell = ws.cell(r + 1, c + 1)
                    value = (
                        cached.cell(r + 1, c + 1).value
                        if cell.data_type == "f"
                        else cell.value
                    )
                    absent = cell.data_type == "f" and value is None
                    if absent:
                        missing.append(cell.coordinate)
                    row.append(Cell(r, c, value, absent))
                rows.append(row)
            sheet = Sheet(ws.title, rows, ws.sheet_state != "visible")
            if missing:
                sheet.warnings.append(
                    f"公式缺少缓存值（{', '.join(missing[:10])}），请用 Excel 重新计算并保存"
                )
            _merged(
                sheet,
                [
                    (m.min_row - 1, m.max_row, m.min_col - 1, m.max_col)
                    for m in ws.merged_cells.ranges
                ],
            )
            result.append(sheet)
        return result
    finally:
        formulas.close()
        if values:
            values.close()


def _xls(data: bytes) -> list[Sheet]:
    book = xlrd.open_workbook(file_contents=data, formatting_info=True, on_demand=True)
    try:
        if book.nsheets > MAX_SHEETS:
            raise FileProblem("工作表数量不能超过 50")
        result, total = [], 0
        for ws in book.sheets():
            total = _shape(ws.nrows, ws.ncols, total)
            rows = []
            for r in range(ws.nrows):
                row = []
                for c in range(ws.ncols):
                    source = ws.cell(r, c)
                    value = source.value
                    if source.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                        if 0 <= source.value < 1:
                            value = value.time()
                    row.append(Cell(r, c, value))
                rows.append(row)
            sheet = Sheet(ws.name, rows, bool(ws.visibility))
            _merged(sheet, ws.merged_cells)
            result.append(sheet)
        return result
    finally:
        book.release_resources()


def _csv(data: bytes) -> list[Sheet]:
    if (
        b"\x00" in data
        or data.startswith(b"PK\x03\x04")
        or data.startswith(b"\xd0\xcf\x11\xe0")
    ):
        raise FileProblem("CSV 内容不像文本，请检查扩展名或转为 UTF-8 编码")
    content = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            content = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if content is None:
        raise FileProblem("无法识别 CSV 编码，请另存为 UTF-8 CSV")
    try:
        dialect = csv.Sniffer().sniff(content[:8192], delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    raw = []
    width = 0
    try:
        for row in csv.reader(
            io.StringIO(content, newline=""), dialect, doublequote=True, strict=True
        ):
            width = max(width, len(row))
            _shape(len(raw) + 1, width, 0)
            raw.append(row)
    except csv.Error as exc:
        raise FileProblem("CSV 分隔或引号格式错误") from exc
    rows = [
        [Cell(r, c, row[c] if c < len(row) else None) for c in range(width)]
        for r, row in enumerate(raw)
    ]
    return [Sheet("CSV", rows)]
